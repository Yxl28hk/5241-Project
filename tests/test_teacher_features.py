"""Offline authorization, document and drafting checks. No live provider calls."""
import io
import time
from unittest.mock import Mock

import pytest
from pptx import Presentation
from pypdf import PdfWriter

from app import create_app
from app.studio import ai_routes
from app.studio.course_ai import CourseAIError, parse_file, validate_result


def fixture_document():
    presentation=Presentation()
    slide=presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text='Explain a Pull Request'
    slide.placeholders[1].text='Explain why you made the change, what changed and what you actually tested. Do not claim unperformed tests passed.'
    buffer=io.BytesIO();presentation.save(buffer)
    return buffer.getvalue()


def fixture_result():
    course=dict(mode='guided',title='First PR',description='Explain a change',category='技术',level='beginner',minutes=15,
                audience='Beginners',objectives='Explain the actual validation',example='Example',task='Write a PR',checklist='Be accurate',reference='Check what you tested',
                sections=[dict(title='PR',body='State the reason and evidence',source_pages=[1],source_quote='Explain a Pull Request')])
    return dict(course=course,review=dict(summary='Check the example',strengths=['Concrete task'],issues=[]),supplements=['Example'])


def test_document_extraction_and_citation_rejection():
    parsed=parse_file(fixture_document(),'pptx')
    assert parsed['total_pages']==1
    data=fixture_result()
    assert validate_result(data,parsed)['sources'][0]['quote_matched']
    data['course']['sections'][0]['source_pages']=[999]
    with pytest.raises(CourseAIError):validate_result(data,parsed)
    writer=PdfWriter();writer.add_blank_page(width=200,height=200)
    buffer=io.BytesIO();writer.write(buffer)
    with pytest.raises(CourseAIError):parse_file(buffer.getvalue(),'pdf')


def test_async_drafting_scoped_csrf_deduplicated_and_cancellable(monkeypatch):
    for key,value in dict(STUDIO_AI_LOCAL_PREVIEW='1',STUDIO_AI_BASE_URL='https://example.invalid',STUDIO_AI_API_KEY='test-placeholder',STUDIO_AI_CONFIG_FILE='').items():monkeypatch.setenv(key,value)
    application=create_app({'TESTING':True,'SECRET_KEY':'test-session','DATABASE_URL':''})
    client=application.test_client();other=application.test_client()
    assert client.get('/api/studio/ai/config').status_code==403
    csrf=client.get('/api/studio/ai/config?preview=1').json['csrf']
    headers={'X-Studio-CSRF':csrf}
    assert client.post('/api/studio/ai/jobs?preview=1').status_code==403
    assert client.post('/api/studio/ai/jobs?preview=1',headers={**headers,'Origin':'https://other.invalid'}).status_code==403
    blob=fixture_document();result=validate_result(fixture_result(),parse_file(blob,'pptx'))
    model=Mock(return_value=result);monkeypatch.setattr(ai_routes,'request_course',model)
    form=lambda:dict(request_id='offline-unique-request',audience='auto',file=(io.BytesIO(blob),'sample.pptx'))
    response=client.post('/api/studio/ai/jobs?preview=1',headers=headers,data=form())
    assert response.status_code==202
    jid=response.json['job_id']
    assert client.post('/api/studio/ai/jobs?preview=1',headers=headers,data=form()).json['job_id']==jid
    assert other.get(f'/api/studio/ai/jobs/{jid}?preview=1').status_code==404
    for _ in range(100):
        job=client.get(f'/api/studio/ai/jobs/{jid}?preview=1').json
        if job['state'] in {'done','error'}:break
        time.sleep(.1)
    assert job['state']=='done',job
    assert 'sample.pptx' in job['result']['course']['source_notes']
    assert not {'owner','token','key','base'} & job.keys()
    model.assert_called_once()
    assert client.delete(f'/api/studio/ai/jobs/{jid}?preview=1',headers=headers).status_code==200
    assert 'result' not in client.get(f'/api/studio/ai/jobs/{jid}?preview=1').json


def test_analytics_and_authoring_require_teacher_and_course_ownership(monkeypatch):
    from app.auth import security
    from app.courses import routes,service,repository
    application=create_app({'TESTING':True,'SECRET_KEY':'test-session','DATABASE_URL':''})
    client=application.test_client()
    assert client.get('/api/courses/7/learners/99').status_code==401
    assert client.post('/api/studio/courses').status_code==401
    with client.session_transaction() as session:session['user_id']=42
    monkeypatch.setattr(security,'get_db',lambda:Mock())
    monkeypatch.setattr(security,'find_role',lambda *_:{'role':'student'})
    assert client.get('/api/courses/7/learners/99').status_code==403
    assert client.post('/api/studio/courses').status_code==403
    monkeypatch.setattr(security,'find_role',lambda *_:{'role':'teacher'})
    monkeypatch.setattr(routes,'get_db',lambda:Mock())
    manager=Mock(return_value=None);monkeypatch.setattr(service,'managed_course',manager)
    assert client.get('/api/courses/7/learners/99').status_code==404
    manager.assert_called_once();assert manager.call_args.args[1:]==(42,7)
    manager.return_value={'id':7,'title':'Scoped course'}
    find=Mock(return_value={'id':99,'name':'Learner','progress':20});monkeypatch.setattr(repository,'learner_detail',find)
    exercises=Mock(return_value=[]);monkeypatch.setattr(repository,'learner_exercises',exercises)
    response=client.get('/api/courses/7/learners/99')
    assert response.status_code==200 and response.json['learner']['id']==99
    assert find.call_args.args[1:]==(7,99) and exercises.call_args.args[1:]==(7,99)
