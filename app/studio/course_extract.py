"""Bounded text-only extraction worker; no network or model credentials."""
import io
import json
import sys
import zipfile

MAX_PAGES = 60
MAX_TEXT = 45000

def extract_document(blob, extension):
    if len(blob)>20*1024*1024:raise ValueError('自动整理最多支持 20MB，请拆分文件后重试。')
    units=[];warnings=[]
    if extension=='pdf':
        if not blob.startswith(b'%PDF'):raise ValueError('这不是有效的 PDF 文件。')
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(blob))
        if reader.is_encrypted:raise ValueError('请先移除 PDF 密码，再上传。')
        total=len(reader.pages)
        if total>MAX_PAGES:raise ValueError('自动整理最多支持 60 页，请按章节拆分后上传。')
        for number,page in enumerate(reader.pages,1):
            try:text=page.extract_text() or ''
            except Exception:text='';warnings.append(f'第 {number} 页文字提取失败，请人工核对。')
            units.append({'page':number,'text':text})
    elif extension=='pptx':
        from pptx import Presentation
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            entries=archive.infolist()
            if len(entries)>4000 or sum(e.file_size for e in entries)>80*1024*1024:
                raise ValueError('演示文稿解压后过大，请压缩图片或拆分文件。')
            if 'ppt/presentation.xml' not in archive.namelist():raise ValueError('这不是有效的 PPTX 文件。')
        presentation=Presentation(io.BytesIO(blob));total=len(presentation.slides)
        if total>MAX_PAGES:raise ValueError('自动整理最多支持 60 张幻灯片，请按章节拆分。')
        def shape_text(shape):
            if shape.has_text_frame:return shape.text
            if shape.has_table:return '\n'.join(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
            if hasattr(shape,'shapes'):return '\n'.join(shape_text(child) for child in shape.shapes)
            return ''
        for number,slide in enumerate(presentation.slides,1):
            units.append({'page':number,'text':'\n'.join(shape_text(s) for s in slide.shapes)})
    else:raise ValueError('自动整理目前支持 PDF 和 PPTX；旧版 PPT 请先另存为 PPTX。')
    total_chars=0;readable=0
    for unit in units:
        unit['text']=unit['text'].replace('\x00','').strip()
        if len(unit['text'])<30:warnings.append(f"第 {unit['page']} 页文字较少或无法提取，可能是封面、图片或扫描页，请检查原资料。")
        if unit['text']:readable+=1
        total_chars+=len(unit['text'])
    if total_chars<80:raise ValueError('未提取到足够文字。扫描件和图片型资料暂不支持 OCR，请上传可复制文字的 PDF/PPTX。')
    if total_chars>MAX_TEXT:raise ValueError('资料文字超过 45000 字符，请拆成更小的章节；本次没有截断或调用模型。')
    warnings.append('仅提取文字及 PPTX 表格文字；图片、图表、公式版式、动画和演讲者备注未被完整理解。')
    return {'pages':units,'total_pages':len(units),'readable_pages':readable,'characters':total_chars,'warnings':warnings}

if __name__=='__main__':
    try:result=extract_document(sys.stdin.buffer.read(20*1024*1024+1),sys.argv[1])
    except ValueError as error:result={'error':str(error)}
    except Exception:result={'error':'文件解析失败，请检查文件是否损坏或改为导出文字型 PDF。'}
    sys.stdout.buffer.write(json.dumps(result,ensure_ascii=False).encode('utf-8'))
